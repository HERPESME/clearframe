"""The clearance report as a Word document.

The same `ClearanceDossier` the HTML renders, laid out as a document somebody
files: a cover page, a document-control block, a numbered schedule of findings,
and a page-numbered footer. It exists because a clearance report is attached to
an E&O application and marked up by counsel, and neither of those happens to a
browser tab.

Real named styles rather than direct formatting, because a Word document whose
headings are only bold text cannot be navigated, cannot generate a table of
contents, and cannot be restyled by whoever receives it.

Everything user-written goes in as TEXT — `add_run` never interprets markup — so
the escaping problem the HTML renderer had cannot exist here. That is a property
of the format, not vigilance, and it is worth saying out loud because the same
values (a production title, a reviewer's decision note) are the ones that
executed as script in the HTML report.
"""

from __future__ import annotations

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from clearframe.dossier import ClearanceDossier, research_is_incomplete
from clearframe.timecode import seconds_to_tc
from clearframe.reportmeta import (
    CONFIDENTIALITY,
    control_rows,
    reference,
    signatories,
    status,
    summary_rows,
)

BAND_RGB = {
    "CRITICAL": RGBColor(0xC0, 0x39, 0x2B),
    "HIGH": RGBColor(0xE6, 0x7E, 0x22),
    "MEDIUM": RGBColor(0x29, 0x80, 0xB9),
    "LOW": RGBColor(0x27, 0xAE, 0x60),
}


def _page_number_footer(section) -> None:
    """`Page N of M` in the footer, as real Word fields.

    Field codes rather than literal text: a printed clearance report is paginated
    by whoever prints it, and a hard-coded number would be wrong the moment
    anyone adds a comment box.
    """
    p = section.footer.paragraphs[0]
    if p.runs:
        # A new section inherits the previous one's footer object, so applying
        # this per section appended the fields a second time and printed
        # "Page  of Page  of ". One footer, written once.
        return
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Page ")
    for instr in ("PAGE", "NUMPAGES"):
        run = p.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        code = OxmlElement("w:instrText")
        code.set(qn("xml:space"), "preserve")
        code.text = f" {instr} "
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(begin)
        run._r.append(code)
        run._r.append(end)
        if instr == "PAGE":
            p.add_run(" of ")


def _styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)


def _kv_table(doc: Document, rows: list[tuple[str, str]], widths=(2.2, 4.0)) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light List Accent 1"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
        cells[0].paragraphs[0].runs[0].bold = True
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def _para(doc: Document, text: str, *, bold=False, italic=False, size=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    return p


def _cover(doc: Document, d: ClearanceDossier) -> None:
    _para(doc, CONFIDENTIALITY, italic=True, size=8)
    doc.add_paragraph()
    doc.add_paragraph()
    heading = doc.add_heading("Rights Clearance Report", level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _para(doc, d.production.title, bold=True, size=18)
    _para(doc, reference(d.production), size=11)
    _para(doc, status(d), italic=True, size=11)
    doc.add_paragraph()
    _kv_table(doc, control_rows(d))
    doc.add_paragraph()

    doc.add_heading("Basis and limitations", level=2)
    _para(doc, d.disclaimer)

    who = signatories(d)
    if who:
        doc.add_heading("Reviewed and signed off by", level=2)
        for name in who:
            doc.add_paragraph(name, style="List Bullet")

    doc.add_section(WD_SECTION.NEW_PAGE)


def _summary(doc: Document, d: ClearanceDossier) -> None:
    doc.add_heading("1. Summary of findings", level=1)
    _kv_table(doc, summary_rows(d), widths=(3.0, 1.2))

    if d.source_work_summary:
        doc.add_heading("1.1 The work itself", level=2)
        _para(doc, str(d.source_work_summary.get("headline", "")))
        if d.source_work_summary.get("action"):
            _para(doc, str(d.source_work_summary["action"]), italic=True)

    if d.cast:
        doc.add_heading("1.2 Credited performers", level=2)
        _para(
            doc,
            (d.cast_summary or {}).get(
                "headline", "Recognised performers, not third-party findings."
            ),
        )
        for credit in d.cast:
            doc.add_paragraph(
                f"{credit.performer}"
                + (f" as {credit.character}" if credit.character else "")
                + (f" — {credit.basis}" if credit.basis else ""),
                style="List Bullet",
            )

    if d.unscanned_ranges:
        doc.add_heading("1.3 Footage not covered by this report", level=2)
        fps = d.production.fps
        for r in d.unscanned_ranges:
            doc.add_paragraph(
                f"{seconds_to_tc(r.start_s, fps=fps)} – {seconds_to_tc(r.end_s, fps=fps)}",
                style="List Bullet",
            )


def _finding(doc: Document, n: int, entry, d: ClearanceDossier) -> None:
    el = entry.element
    doc.add_heading(f"2.{n} {el.label}", level=2)

    band = doc.add_paragraph()
    run = band.add_run(f"{entry.risk.band.value} · score {entry.risk.score}")
    run.bold = True
    run.font.color.rgb = BAND_RGB.get(entry.risk.band.value, RGBColor(0, 0, 0))
    band.add_run(f"    {el.category.value}")
    if entry.risk.de_minimis:
        band.add_run("    DE MINIMIS")

    fps = d.production.fps
    if el.timing_reliable:
        times = "; ".join(
            f"{seconds_to_tc(r.start_s, fps=fps)}–{seconds_to_tc(r.end_s, fps=fps)}"
            for r in el.time_ranges
        )
        _para(doc, f"Appears: {times}")
    else:
        _para(doc, f"Timecodes unusable. {el.timing_note}", italic=True)
    if el.description:
        _para(doc, el.description)

    if entry.route is not None:
        doc.add_heading("How this was resolved", level=3)
        _para(doc, entry.route.rationale)
        if entry.route.basis:
            _para(doc, f"Authority: {entry.route.basis}")
        if entry.route.disposition:
            _para(doc, f"Required action: {entry.route.disposition}", bold=True)

    doc.add_heading("Rights research", level=3)
    if entry.research is not None and not research_is_incomplete(entry.research):
        r = entry.research
        _kv_table(
            doc,
            [
                ("Rights holder", r.owner or "not established"),
                ("Confidence", r.owner_confidence),
                ("Licensing posture", r.licensing_posture.value),
                ("Contact", r.licensing_contact or "—"),
                ("Estimated licence", r.estimated_license_cost_band or "—"),
            ],
        )
        if r.litigation_history:
            _para(doc, "Enforcement history:", bold=True)
            for item in r.litigation_history:
                doc.add_paragraph(item, style="List Bullet")
        if r.basis:
            _para(doc, "Evidence basis:", bold=True)
            for b in r.basis:
                doc.add_paragraph(
                    f"[{b.field}] {b.url} — “{b.excerpt}” "
                    f"({b.confidence} confidence: {b.reasoning})",
                    style="List Bullet",
                )
    else:
        _para(
            doc,
            "RESEARCH INCOMPLETE — no rights holder was established for this "
            "element. It must be cleared or removed before delivery.",
            italic=True,
        )

    if entry.coverage is not None:
        doc.add_heading("Rights already held", level=3)
        _para(doc, f"{entry.coverage.status.value} — {entry.coverage.note}")
        for gap in entry.coverage.gaps:
            doc.add_paragraph(gap, style="List Bullet")

    if entry.corroboration is not None:
        doc.add_heading("Identity verification", level=3)
        _para(doc, f"{entry.corroboration.verdict.value} — {entry.corroboration.note}")

    rows = [
        (
            t.territory,
            f"{t.band.value} — {t.rationale}"
            + (f" [{t.authority}]" if t.authority else ""),
        )
        for t in (entry.territory or [])
    ]
    if rows:
        doc.add_heading("Territory exposure", level=3)
        _kv_table(doc, rows, widths=(1.2, 5.0))

    if entry.freshness:
        doc.add_heading("Live rights-holder signals", level=3)
        for signal in entry.freshness:
            prefix = "ENFORCEMENT SIGNAL — " if signal.material else ""
            doc.add_paragraph(
                f"{prefix}{signal.title} ({signal.url}) — {signal.excerpt}",
                style="List Bullet",
            )

    if entry.options:
        doc.add_heading("Remediation options", level=3)
        for option in entry.options:
            doc.add_paragraph(
                f"{option.kind}: {option.summary}"
                + (f" [{option.est_cost_band}]" if option.est_cost_band else ""),
                style="List Bullet",
            )

    doc.add_heading("Decision", level=3)
    if entry.decision is not None:
        dec = entry.decision
        _para(
            doc,
            f"{dec.action} — {dec.note or 'no note recorded'} "
            f"({dec.reviewer}, {dec.role})",
        )
    else:
        _para(doc, "PENDING REVIEW", bold=True)


def _audit(doc: Document, d: ClearanceDossier) -> None:
    if not d.audit:
        return
    doc.add_heading("3. Audit trail", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light List Accent 1"
    for cell, name in zip(
        table.rows[0].cells, ("When", "Actor", "Role", "Event", "Detail")
    ):
        cell.text = name
        cell.paragraphs[0].runs[0].bold = True
    for event in d.audit:
        cells = table.add_row().cells
        for cell, value in zip(
            cells, (event.at, event.actor, event.role, event.event, event.detail)
        ):
            cell.text = value


def render_dossier_docx(d: ClearanceDossier, path) -> None:
    """Write the report to `path` as a .docx.

    `path` may be a filesystem path or an open binary stream — python-docx takes
    either, and the tests render into memory rather than to a temp file.
    """
    doc = Document()
    _styles(doc)
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.9)
        section.top_margin = section.bottom_margin = Inches(0.8)

    _cover(doc, d)
    _summary(doc, d)

    doc.add_heading("2. Schedule of findings", level=1)
    if not d.entries:
        _para(doc, "No clearance findings were identified in this footage.")
    for n, entry in enumerate(d.entries, start=1):
        _finding(doc, n, entry, d)

    _audit(doc, d)
    for section in doc.sections:
        _page_number_footer(section)
    doc.save(path if hasattr(path, "write") else str(path))
